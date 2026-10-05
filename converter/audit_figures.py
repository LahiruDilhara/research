import docx, os, re
from PIL import Image
from pathlib import Path

WORKSPACE = Path('.').resolve()
FIGURES_DIR = WORKSPACE / 'figures'

print('=== 1. AUDIT OF ALL FIGURES IN LATEX CHAPTER FILES ===')
tex_figs = []
for p in sorted(WORKSPACE.glob('chapters/*.tex')):
    content = p.read_text(encoding='utf-8')
    matches = re.finditer(r'\\begin\{figure\}.*?\\caption\{(.*?)\}.*?\\label\{(.*?)\}.*?\\end\{figure\}', content, re.DOTALL)
    for m in matches:
        cap = m.group(1).strip()
        lbl = m.group(2).strip()
        block = m.group(0)
        
        imgs = re.findall(r'\\includegraphics(?:\[.*?\])?\{(.*?)\}', block)
        has_tikz = 'tikzpicture' in block
        
        tex_figs.append({
            'file': p.name,
            'label': lbl,
            'caption': cap,
            'has_tikz': has_tikz,
            'imgs': imgs
        })

print(f'Total figures found in LaTeX source files: {len(tex_figs)}')
for idx, f in enumerate(tex_figs, 1):
    src = f['imgs'] if f['imgs'] else ('[TikZ Diagram]' if f['has_tikz'] else '[Unknown]')
    print(f'{idx:2d}. Label: {f["label"]}')
    print(f'    File: {f["file"]}')
    print(f'    Caption: {f["caption"]}')
    print(f'    Source: {src}')

print('\n=== 2. AUDIT OF ALL FIGURE IMAGES ON DISK ===')
all_pngs = sorted(FIGURES_DIR.glob('*.*'))
for img_path in all_pngs:
    try:
        with Image.open(img_path) as im:
            print(f'  {img_path.name:55s} | Size: {im.size[0]:4d}x{im.size[1]:4d} px | Format: {im.format:4s} | Mode: {im.mode}')
    except Exception as e:
        print(f'  {img_path.name:55s} | ERROR: {e}')

print('\n=== 3. AUDIT OF FIGURES IN THESIS.DOCX ===')
doc = docx.Document('output/thesis.docx')

docx_figs = []
for p_idx, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    # Check if this is a body figure caption
    if t.startswith('Figure ') and ('Figure ' in t[:10]) and (':' in t[:15]):
        drawing_details = []
        for back_idx in range(max(0, p_idx-4), p_idx):
            back_p = doc.paragraphs[back_idx]
            for blip in back_p._p.xpath('.//a:blip'):
                rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if rId:
                    rel = doc.part.rels.get(rId)
                    target = rel.target_ref if rel else 'UNKNOWN'
                    drawing_details.append(target)
                    
        docx_figs.append({
            'caption': t,
            'drawings': drawing_details,
            'p_idx': p_idx
        })

print(f'Total figure captions in DOCX body: {len(docx_figs)}')
for idx, df in enumerate(docx_figs, 1):
    print(f'{idx:2d}. Caption: {df["caption"]}')
    print(f'    Embed Targets: {df["drawings"]}')
